"""Read a reviewed work schedule from the field dashboard Excel template."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from hashlib import sha256
from io import BytesIO
import re

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


HEADERS = (
    "작업ID", "작업일", "시작", "종료", "동/구역", "세부 위치", "공종", "작업 내용",
    "주요 장비", "인원(명)", "협력업체", "작업책임자", "계획된 안전조치", "추가 확인",
    "검토 상태", "변경 사항",
)
REQUIRED = {"작업일", "시작", "종료", "동/구역", "작업 내용"}
HEADER_ALIASES = {
    "작업ID": ("작업ID", "작업번호", "공정번호", "관리번호"),
    "작업일": ("작업일", "작업일자", "일자", "날짜"),
    "시작": ("시작", "시작시간", "착수시간", "작업시작"),
    "종료": ("종료", "종료시간", "완료시간", "작업종료"),
    "동/구역": ("동/구역", "구역", "작업구역", "작업장소", "장소"),
    "세부 위치": ("세부위치", "위치", "작업위치"),
    "공종": ("공종", "공정", "작업공종"),
    "작업 내용": ("작업내용", "작업명", "작업", "작업계획"),
    "주요 장비": ("주요장비", "장비", "사용장비"),
    "인원(명)": ("인원(명)", "인원", "작업인원", "투입인원"),
    "협력업체": ("협력업체", "업체", "시공사"),
    "작업책임자": ("작업책임자", "책임자", "담당자", "작업담당자"),
    "계획된 안전조치": ("계획된안전조치", "안전조치", "안전대책", "예방조치"),
    "추가 확인": ("추가확인", "확인사항", "작업전확인", "주의사항"),
    "검토 상태": ("검토상태", "검토결과"),
    "변경 사항": ("변경사항", "변경내용"),
}


def _normalize_header(value: str) -> str:
    return "".join(value.lower().split())


ALIAS_TO_HEADER = {
    _normalize_header(alias): canonical
    for canonical, aliases in HEADER_ALIASES.items()
    for alias in aliases
}


@dataclass(frozen=True)
class WorkItem:
    work_id: str
    day: date
    start: time
    end: time
    area: str
    location: str
    trade: str
    activity: str
    equipment: str
    people: int | None
    contractor: str
    owner: str
    planned_controls: str
    follow_up: str
    review_status: str
    change_note: str
    sheet: str
    row: int


@dataclass(frozen=True)
class NoWorkConfirmation:
    day: date
    confirmed_by: str
    note: str
    at: datetime
    revision_token: str


@dataclass(frozen=True)
class WorkPlan:
    site: str
    items: tuple[WorkItem, ...]
    issues: tuple[str, ...]
    no_work_confirmations: tuple[NoWorkConfirmation, ...] = ()


@dataclass(frozen=True)
class PlanChanges:
    added: tuple[str, ...]
    removed: tuple[str, ...]
    changed: tuple[str, ...]
    unchanged: tuple[str, ...]


def compare_plans(previous: WorkPlan, updated: WorkPlan) -> PlanChanges:
    """Compare source content by work ID, ignoring its location within the workbook."""
    old = {item.work_id: item for item in previous.items}
    new = {item.work_id: item for item in updated.items}
    def content(item: WorkItem) -> tuple:
        return (
            item.day, item.start, item.end, item.area, item.location, item.trade,
            item.activity, item.equipment, item.people, item.contractor, item.owner,
            item.planned_controls, item.follow_up, item.review_status, item.change_note,
        )
    shared = old.keys() & new.keys()
    return PlanChanges(
        tuple(sorted(new.keys() - old.keys())),
        tuple(sorted(old.keys() - new.keys())),
        tuple(sorted(work_id for work_id in shared if content(old[work_id]) != content(new[work_id]))),
        tuple(sorted(work_id for work_id in shared if content(old[work_id]) == content(new[work_id]))),
    )


def _string(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _day(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        digits = str(int(value))
        if value == int(value) and len(digits) == 8:
            return date.fromisoformat(f"{digits[:4]}-{digits[4:6]}-{digits[6:]}")
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.date()
    if isinstance(value, str):
        normalized = re.sub(r"\s*[-./년월]\s*", "-", value.strip())
        normalized = re.sub(r"\s*일\s*$", "", normalized).rstrip("-")
        match = re.search(r"\d{4}-\d{1,2}-\d{1,2}", normalized)
        if match:
            year, month, day = (int(part) for part in match.group().split("-"))
            return date(year, month, day)
    raise ValueError("작업일 형식을 확인하세요")


def _time(value: object) -> time:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, timedelta):
        seconds = int(value.total_seconds())
        if 0 <= seconds < 86400:
            return (datetime.min + value).time()
    if isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = round(value * 86400)
        if seconds < 86400:
            return (datetime.min + timedelta(seconds=seconds)).time()
    if isinstance(value, (int, float)) and value == int(value):
        number = int(value)
        if 0 <= number <= 23:
            return time(number)
        if 100 <= number <= 2359 and number % 100 < 60 and number // 100 < 24:
            return time(number // 100, number % 100)
    if isinstance(value, str):
        text = value.strip()
        meridiem = None
        if text.startswith("오전 ") or text.startswith("오후 "):
            meridiem = text.startswith("오후 ")
            text = text[3:].strip()
        match = re.fullmatch(r"(\d{1,2})시(?:\s*(\d{1,2})분?)?", text)
        if match:
            hour = int(match.group(1))
            return time(hour % 12 + (12 if meridiem else 0) if meridiem is not None else hour, int(match.group(2) or 0))
        match = re.fullmatch(r"(\d{1,2})(?::(\d{1,2}))?", text)
        if match:
            hour = int(match.group(1))
            return time(hour % 12 + (12 if meridiem else 0) if meridiem is not None else hour, int(match.group(2) or 0))
        return time.fromisoformat(text)
    raise ValueError("시작·종료 시간 형식을 확인하세요")


def read_work_plan(content: bytes) -> WorkPlan:
    """Read the template; skip invalid rows and report their exact source location."""
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("파일 크기는 10MB 이하여야 합니다.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("엑셀 파일을 읽을 수 없습니다. .xlsx 형식을 확인하세요.") from exc

    items: list[WorkItem] = []
    issues: list[str] = []
    seen: set[str] = set()
    site = "현장명 미입력"
    found_table = False
    scanned_rows = 0
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header: dict[str, int] | None = None
            for row_number, values in enumerate(rows, start=1):
                scanned_rows += 1
                if scanned_rows > 10_000:
                    raise ValueError("작업계획서는 전체 10,000행 이하여야 합니다.")
                cells = [_string(value) for value in values]
                if cells and _normalize_header(cells[0]) == "현장명" and len(cells) > 1:
                    site = cells[1] or site
                if header is None:
                    candidate = {}
                    for index, name in enumerate(cells):
                        canonical = ALIAS_TO_HEADER.get(_normalize_header(name))
                        if canonical and canonical not in candidate:
                            candidate[canonical] = index
                    if REQUIRED.issubset(candidate):
                        header = candidate
                        found_table = True
                    continue
                if not any(value is not None and _string(value) for value in values):
                    continue
                def cell(name: str) -> object:
                    index = header.get(name)
                    return values[index] if index is not None and index < len(values) else None

                if not any(_string(cell(name)) for name in REQUIRED):
                    # Notes below the table are not work rows.
                    continue
                work_id = _string(cell("작업ID"))
                source = f"{sheet.title} {row_number}행"
                if work_id and work_id in seen:
                    issues.append(f"{source}: 작업ID {work_id}가 중복되어 제외했습니다.")
                    continue
                try:
                    day = _day(cell("작업일"))
                    start = _time(cell("시작"))
                    end = _time(cell("종료"))
                    if end <= start:
                        raise ValueError("종료 시간은 시작 시간보다 늦어야 합니다")
                    area = _string(cell("동/구역"))
                    activity = _string(cell("작업 내용"))
                    if not area or not activity:
                        raise ValueError("동/구역과 작업 내용을 입력하세요")
                    text_values = {
                        name: _string(cell(name))
                        for name in (
                            "세부 위치", "공종", "주요 장비", "협력업체", "작업책임자",
                            "계획된 안전조치", "추가 확인", "검토 상태", "변경 사항",
                        )
                    }
                    text_limits = {
                        "작업ID": 100, "동/구역": 150, "작업 내용": 300,
                        "세부 위치": 300, "공종": 120, "주요 장비": 200,
                        "협력업체": 150, "작업책임자": 100,
                        "계획된 안전조치": 2000, "추가 확인": 2000,
                        "검토 상태": 100, "변경 사항": 2000,
                    }
                    text_values.update({"동/구역": area, "작업 내용": activity, "작업ID": work_id})
                    too_long = next((name for name, value in text_values.items() if len(value) > text_limits[name]), None)
                    if too_long:
                        raise ValueError(f"{too_long}이 입력 한도를 넘었습니다")
                    people_value = cell("인원(명)")
                    if people_value in (None, ""):
                        people = None
                    else:
                        people = int(people_value)
                        if str(people) != str(people_value).strip() and float(people_value) != people:
                            raise ValueError("인원은 정수여야 합니다")
                    if people is not None and people < 0:
                        raise ValueError("인원은 0 이상이어야 합니다")
                    if not work_id:
                        identity = f"{day.isoformat()}|{start.isoformat()}|{end.isoformat()}|{area}|{activity}"
                        work_id = "AUTO-" + sha256(identity.encode("utf-8")).hexdigest()[:10].upper()
                    if work_id in seen:
                        raise ValueError(f"동일한 작업이 중복되었습니다 ({work_id})")
                except (ValueError, TypeError, OverflowError) as exc:
                    issues.append(f"{source}: {exc}. 이 작업은 제외했습니다.")
                    continue
                if len(items) >= 500:
                    raise ValueError("작업계획서 한 건은 최대 500개 작업까지 지원합니다.")
                seen.add(work_id)
                items.append(WorkItem(
                    work_id, day, start, end, area, text_values["세부 위치"],
                    text_values["공종"], activity, text_values["주요 장비"], people,
                    text_values["협력업체"], text_values["작업책임자"],
                    text_values["계획된 안전조치"], text_values["추가 확인"],
                    text_values["검토 상태"], text_values["변경 사항"],
                    sheet.title, row_number,
                ))
    finally:
        workbook.close()
    if not found_table:
        raise ValueError("작업표를 찾지 못했습니다. 샘플 계획서의 열 이름을 확인하세요.")
    if len(site) > 200:
        raise ValueError("현장명은 200자 이내로 입력하세요.")
    if not items:
        raise ValueError("사용 가능한 작업이 없습니다. 날짜·시간·작업 내용을 확인하세요.")
    return WorkPlan(site, tuple(sorted(items, key=lambda item: (item.day, item.start, item.work_id))), tuple(issues))


def overlapping_pairs(items: tuple[WorkItem, ...], day: date) -> tuple[tuple[WorkItem, WorkItem], ...]:
    """Flag potential coordination needs in the same broad area; no hazard verdict."""
    daily = [item for item in items if item.day == day]
    return tuple(
        (left, right)
        for index, left in enumerate(daily)
        for right in daily[index + 1:]
        if left.area == right.area and left.start < right.end and right.start < left.end
    )
